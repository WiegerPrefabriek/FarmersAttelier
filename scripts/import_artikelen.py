#!/usr/bin/env python3
"""Leest 'Farmers Atelier Artikelen.xlsx' en zet er nette JSON van.

De Excel is met de hand bijgehouden en is daar ook naar: de maatkolom heet op het
ene tabblad "Maat:" en op het andere "Size:", XXL en 2XL staan door elkaar, en
alleen bij de shirts is de voorraad ingevuld. Dit script maakt daar één vorm van
zodat het dashboard er niet omheen hoeft te werken.

    ./.venv/bin/python scripts/import_artikelen.py "<pad naar xlsx>"

Schrijft data/bronnen/artikelen.json en data/bronnen/retourscans.json.
"""

from __future__ import annotations

import json
import os
import re
import sys
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import config  # noqa: E402

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl ontbreekt: ./.venv/bin/pip install openpyxl")

# De tabbladen met artikelen, en de groep waaronder ze in het dashboard vallen.
ARTIKELBLADEN = {
    "T-shirts & longsleeves": "Shirts",
    "Hoodies & knits": "Truien",
    "Accesoires": "Accessoires",
}
RETOURBLAD = "Gescande retouren"

# Eén vaste vololgorde, zodat maten in het dashboard niet alfabetisch komen te staan
# (dan staat L vóór M en XS achter XL, wat nergens op slaat).
MAATVOLGORDE = ["XS", "S", "M", "L", "XL", "2XL", "3XL", "One size"]
MAAT_ALIAS = {"XXL": "2XL", "XXXL": "3XL", "ONE": "One size", "ONESIZE": "One size"}


def normaliseer_maat(ruw: str | None) -> str | None:
    """'Maat: XXL' en 'Size: 2XL' zijn dezelfde maat. Geeft None bij niets."""
    if not ruw:
        return None
    tekst = str(ruw).strip()
    # "Maat: XL" / "Size: XL" → "XL"
    tekst = re.sub(r"^\s*(maat|size)\s*:\s*", "", tekst, flags=re.I).strip()
    if not tekst:
        return None
    sleutel = tekst.upper().replace(" ", "")
    return MAAT_ALIAS.get(sleutel, tekst.upper() if len(tekst) <= 3 else tekst)


def splits_naam(naam: str) -> tuple[str, str | None]:
    """'Hayline long sleeve | Black' → ('Hayline long sleeve', 'Black')."""
    if "|" in naam:
        model, _, kleur = naam.partition("|")
        model, kleur = model.strip(), kleur.strip()
        return model, (kleur or None)
    return naam.strip(), None


def kolomindex(kop: list, *woorden: str) -> int | None:
    """Zoekt de eerste kolom waarvan de kop een van de woorden bevat."""
    for i, k in enumerate(kop):
        if not k:
            continue
        laag = str(k).lower()
        if any(w in laag for w in woorden):
            return i
    return None


def lees_artikelen(wb) -> list[dict]:
    regels = []
    for blad, groep in ARTIKELBLADEN.items():
        if blad not in wb.sheetnames:
            print(f"  ! tabblad {blad!r} ontbreekt, overgeslagen")
            continue
        ws = wb[blad]
        kop = [c.value for c in ws[1]]
        i_naam = kolomindex(kop, "productnaam") or 0
        i_code = kolomindex(kop, "productcode")
        i_attr = kolomindex(kop, "attribut")
        i_voorraad = kolomindex(kop, "voorraad")
        aantal = 0
        for rij in ws.iter_rows(min_row=2, values_only=True):
            naam = rij[i_naam] if i_naam < len(rij) else None
            if not naam:
                continue
            model, kleur = splits_naam(str(naam))
            voorraad = None
            if i_voorraad is not None and i_voorraad < len(rij):
                waarde = rij[i_voorraad]
                if isinstance(waarde, (int, float)):
                    voorraad = int(waarde)
            regels.append({
                "groep": groep,
                "blad": blad,
                "model": model,
                "kleur": kleur,
                "maat": normaliseer_maat(rij[i_attr] if i_attr is not None and i_attr < len(rij) else None),
                "sku": (str(rij[i_code]).strip() if i_code is not None and i_code < len(rij) and rij[i_code] else None),
                "voorraad": voorraad,
            })
            aantal += 1
        print(f"  {blad}: {aantal} varianten")
    return regels


def lees_retourscans(wb) -> list[str]:
    if RETOURBLAD not in wb.sheetnames:
        return []
    ws = wb[RETOURBLAD]
    codes, gezien = [], set()
    for rij in ws.iter_rows(min_row=1, values_only=True):
        if not rij or not rij[0]:
            continue
        code = str(rij[0]).strip().rstrip("\\")  # één regel had een losse backslash
        if code and code not in gezien:
            gezien.add(code)
            codes.append(code)
    return codes


def vervoerder(code: str) -> str:
    """Leidt de vervoerder af uit de vorm van het trackingnummer."""
    if re.match(r"^3S(DFC|WLT)", code):
        return "DHL"
    if code.startswith("3S"):
        return "PostNL"
    if code.startswith("JVGL"):
        return "DHL eCommerce"
    if code.startswith("1Z"):
        return "UPS"
    if code.startswith("%"):
        return "streepjescode"
    if re.match(r"^[A-Z]{2}\d+[A-Z]{2}$", code):
        return "internationaal"
    return "onbekend"


def main() -> int:
    pad = sys.argv[1] if len(sys.argv) > 1 else os.path.expanduser(
        "~/Downloads/Farmers Atelier Artikelen.xlsx")
    if not os.path.exists(pad):
        return print(f"Bestand niet gevonden: {pad}") or 1

    print(f"Lezen: {pad}")
    wb = openpyxl.load_workbook(pad, data_only=True)

    artikelen = lees_artikelen(wb)
    scans = lees_retourscans(wb)
    print(f"  {RETOURBLAD}: {len(scans)} unieke retourzendingen")

    uit = os.path.join(config.DATA_DIR, "bronnen")
    os.makedirs(uit, exist_ok=True)

    with open(os.path.join(uit, "artikelen.json"), "w", encoding="utf-8") as f:
        json.dump({"bron": os.path.basename(pad), "regels": artikelen}, f,
                  ensure_ascii=False, indent=1)

    per_vervoerder: dict[str, int] = defaultdict(int)
    for c in scans:
        per_vervoerder[vervoerder(c)] += 1
    with open(os.path.join(uit, "retourscans.json"), "w", encoding="utf-8") as f:
        json.dump({"bron": os.path.basename(pad),
                   "aantal": len(scans),
                   "per_vervoerder": dict(per_vervoerder),
                   "codes": [{"tracking": c, "vervoerder": vervoerder(c)} for c in scans]},
                  f, ensure_ascii=False, indent=1)

    met = [r for r in artikelen if r["voorraad"] is not None]
    totaal = sum(r["voorraad"] for r in met)
    print(f"\nKlaar. {len(artikelen)} varianten, waarvan {len(met)} met een voorraadgetal.")
    print(f"Voorraad opgeteld: {totaal} stuks.")
    print(f"Retourzendingen gescand: {len(scans)}.")
    print(f"Geschreven naar {uit}/")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
